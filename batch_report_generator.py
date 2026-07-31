from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# =========================================================
# 样式配置
# =========================================================

TITLE_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)

PASS_FILL = PatternFill(
    "solid",
    fgColor="E2F0D9",
)

FAIL_FILL = PatternFill(
    "solid",
    fgColor="F4CCCC",
)

UNKNOWN_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
)

HEADER_FONT = Font(
    bold=True,
    color="1F1F1F",
)

TITLE_FONT = Font(
    bold=True,
    color="FFFFFF",
    size=16,
)

THIN_BORDER = Border(
    left=Side(
        style="thin",
        color="B7B7B7",
    ),
    right=Side(
        style="thin",
        color="B7B7B7",
    ),
    top=Side(
        style="thin",
        color="B7B7B7",
    ),
    bottom=Side(
        style="thin",
        color="B7B7B7",
    ),
)


# =========================================================
# 基础辅助函数
# =========================================================

def safe_text(
    value: object,
) -> str:
    """
    将任意值转换为Excel可写入的字符串。
    """
    if value is None:
        return ""

    return str(value)


def apply_result_fill(
    cell,
    result: str,
) -> None:
    """
    根据PASS、FAIL和无法判断设置单元格底色。
    """
    normalized_result = safe_text(
        result
    ).strip()

    if normalized_result == "PASS":
        cell.fill = PASS_FILL

    elif normalized_result == "FAIL":
        cell.fill = FAIL_FILL

    else:
        cell.fill = UNKNOWN_FILL


def style_header_row(
    worksheet,
    row_number: int,
) -> None:
    """
    设置表头行样式。
    """
    for cell in worksheet[
        row_number
    ]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = THIN_BORDER


def style_data_range(
    worksheet,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
) -> None:
    """
    为数据区域增加边框、自动换行和垂直居中。
    """
    if max_row < min_row:
        return

    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def set_column_widths(
    worksheet,
    widths: dict[int, float],
) -> None:
    """
    设置列宽。
    """
    for column_index, width in widths.items():
        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = width


def write_title(
    worksheet,
    title: str,
    end_column: int,
) -> None:
    """
    在工作表第一行写入合并标题。
    """
    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    title_cell = worksheet.cell(
        row=1,
        column=1,
        value=title,
    )

    title_cell.fill = TITLE_FILL
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 28


# =========================================================
# 批次统计
# =========================================================

def calculate_batch_counts(
    label_results: list[dict[str, Any]],
) -> tuple[int, int, int]:
    """
    统计标签与图纸比对的PASS、FAIL和无法判断数量。
    """
    pass_count = 0
    fail_count = 0
    unknown_count = 0

    for item in label_results:
        result = safe_text(
            item.get(
                "图纸比对结果",
                "",
            )
        ).strip()

        if result == "PASS":
            pass_count += 1

        elif result == "FAIL":
            fail_count += 1

        else:
            unknown_count += 1

    return (
        pass_count,
        fail_count,
        unknown_count,
    )


# =========================================================
# 工作表：批次总览
# =========================================================

def build_batch_overview_sheet(
    workbook: Workbook,
    drawing_filename: str,
    batch_overall_result: str,
    label_results: list[dict[str, Any]],
    batch_summary_rows: list[dict[str, Any]],
    inconsistent_rows: list[dict[str, Any]],
    detection_time: str,
) -> None:
    """
    创建“批次总览”工作表。
    """
    worksheet = workbook.active
    worksheet.title = "批次总览"

    write_title(
        worksheet,
        "图纸与多标签批量一致性检测报告",
        8,
    )

    (
        label_pass_count,
        label_fail_count,
        label_unknown_count,
    ) = calculate_batch_counts(
        label_results
    )

    overview_rows = [
        [
            "检测时间",
            detection_time,
            "图纸文件名",
            drawing_filename,
        ],
        [
            "标签数量",
            len(label_results),
            "批次最终结果",
            batch_overall_result,
        ],
        [
            "与图纸一致标签数",
            label_pass_count,
            "与图纸异常标签数",
            label_fail_count,
        ],
        [
            "无法判断标签数",
            label_unknown_count,
            "横向异常字段数",
            len(inconsistent_rows),
        ],
    ]

    start_row = 3

    for row_offset, row_values in enumerate(
        overview_rows
    ):
        current_row = start_row + row_offset

        for column_index, value in enumerate(
            row_values,
            start=1,
        ):
            cell = worksheet.cell(
                row=current_row,
                column=column_index,
                value=value,
            )

            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if column_index in (
                1,
                3,
            ):
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL

    result_cell = worksheet.cell(
        row=4,
        column=4,
    )

    apply_result_fill(
        result_cell,
        batch_overall_result,
    )

    summary_start_row = 9

    worksheet.cell(
        row=summary_start_row,
        column=1,
        value="各标签检测汇总",
    ).font = Font(
        bold=True,
        size=13,
    )

    summary_headers = [
        "标签序号",
        "标签文件名",
        "是否基准",
        "一致字段数量",
        "异常字段数量",
        "无法判断数量",
        "标签检测结果",
    ]

    header_row = summary_start_row + 1

    for column_index, header in enumerate(
        summary_headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
    )

    for row_index, summary_row in enumerate(
        batch_summary_rows,
        start=header_row + 1,
    ):
        for column_index, header in enumerate(
            summary_headers,
            start=1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=summary_row.get(
                    header,
                    "",
                ),
            )

        apply_result_fill(
            worksheet.cell(
                row=row_index,
                column=7,
            ),
            safe_text(
                summary_row.get(
                    "标签检测结果",
                    "",
                )
            ),
        )

    style_data_range(
        worksheet,
        header_row + 1,
        header_row + len(batch_summary_rows),
        1,
        len(summary_headers),
    )

    set_column_widths(
        worksheet,
        {
            1: 14,
            2: 38,
            3: 14,
            4: 16,
            5: 16,
            6: 16,
            7: 18,
            8: 18,
        },
    )

    worksheet.freeze_panes = "A10"


# =========================================================
# 工作表：标签与图纸
# =========================================================

def build_drawing_comparison_sheet(
    workbook: Workbook,
    label_results: list[dict[str, Any]],
) -> None:
    """
    创建“标签与图纸”工作表。
    """
    worksheet = workbook.create_sheet(
        "标签与图纸"
    )

    headers = [
        "标签序号",
        "标签文件名",
        "字段名称",
        "图纸值",
        "标签值",
        "检测结果",
        "异常说明",
    ]

    write_title(
        worksheet,
        "每张标签与图纸逐字段比对",
        len(headers),
    )

    header_row = 3

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
    )

    current_row = header_row + 1

    for label_index, item in enumerate(
        label_results,
        start=1,
    ):
        filename = safe_text(
            item.get(
                "文件名",
                f"标签{label_index}",
            )
        )

        comparison_rows = item.get(
            "图纸比对明细",
            [],
        )

        for comparison_row in comparison_rows:
            values = [
                label_index,
                filename,
                comparison_row.get(
                    "字段名称",
                    "",
                ),
                comparison_row.get(
                    "图纸值",
                    "",
                ),
                comparison_row.get(
                    "标签值",
                    "",
                ),
                comparison_row.get(
                    "检测结果",
                    "",
                ),
                comparison_row.get(
                    "异常说明",
                    "",
                ),
            ]

            for column_index, value in enumerate(
                values,
                start=1,
            ):
                worksheet.cell(
                    row=current_row,
                    column=column_index,
                    value=value,
                )

            apply_result_fill(
                worksheet.cell(
                    row=current_row,
                    column=6,
                ),
                safe_text(
                    comparison_row.get(
                        "检测结果",
                        "",
                    )
                ),
            )

            current_row += 1

    style_data_range(
        worksheet,
        header_row + 1,
        current_row - 1,
        1,
        len(headers),
    )

    set_column_widths(
        worksheet,
        {
            1: 12,
            2: 36,
            3: 18,
            4: 24,
            5: 24,
            6: 16,
            7: 42,
        },
    )

    worksheet.freeze_panes = "A4"


# =========================================================
# 工作表：标签横向对比
# =========================================================

def build_label_horizontal_sheet(
    workbook: Workbook,
    batch_detail_rows: list[dict[str, Any]],
) -> None:
    """
    创建“标签横向对比”工作表。
    """
    worksheet = workbook.create_sheet(
        "标签横向对比"
    )

    headers = [
        "标签序号",
        "字段名称",
        "基准标签",
        "基准值",
        "对比标签",
        "对比值",
        "检测结果",
        "异常说明",
    ]

    write_title(
        worksheet,
        "以第一张标签为基准的横向逐字段对比",
        len(headers),
    )

    header_row = 3

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
    )

    for row_index, detail_row in enumerate(
        batch_detail_rows,
        start=header_row + 1,
    ):
        for column_index, header in enumerate(
            headers,
            start=1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=detail_row.get(
                    header,
                    "",
                ),
            )

        apply_result_fill(
            worksheet.cell(
                row=row_index,
                column=7,
            ),
            safe_text(
                detail_row.get(
                    "检测结果",
                    "",
                )
            ),
        )

    style_data_range(
        worksheet,
        header_row + 1,
        header_row + len(batch_detail_rows),
        1,
        len(headers),
    )

    set_column_widths(
        worksheet,
        {
            1: 12,
            2: 18,
            3: 36,
            4: 24,
            5: 36,
            6: 24,
            7: 16,
            8: 42,
        },
    )

    worksheet.freeze_panes = "A4"


# =========================================================
# 工作表：异常字段汇总
# =========================================================

def build_inconsistent_sheet(
    workbook: Workbook,
    inconsistent_rows: list[dict[str, Any]],
) -> None:
    """
    创建“不一致字段汇总”工作表。
    """
    worksheet = workbook.create_sheet(
        "不一致字段汇总"
    )

    if inconsistent_rows:
        dynamic_headers: list[str] = []

        for row in inconsistent_rows:
            for key in row.keys():
                if key not in dynamic_headers:
                    dynamic_headers.append(
                        key
                    )

    else:
        dynamic_headers = [
            "字段名称",
            "异常类型",
        ]

    write_title(
        worksheet,
        "多标签不一致字段汇总",
        max(
            len(dynamic_headers),
            2,
        ),
    )

    header_row = 3

    for column_index, header in enumerate(
        dynamic_headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
    )

    if inconsistent_rows:
        for row_index, inconsistent_row in enumerate(
            inconsistent_rows,
            start=header_row + 1,
        ):
            for column_index, header in enumerate(
                dynamic_headers,
                start=1,
            ):
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=inconsistent_row.get(
                        header,
                        "",
                    ),
                )

    else:
        worksheet.cell(
            row=header_row + 1,
            column=1,
            value="未发现标签之间存在字段差异。",
        )

        worksheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=max(
                len(dynamic_headers),
                2,
            ),
        )

        worksheet.cell(
            row=header_row + 1,
            column=1,
        ).fill = PASS_FILL

    style_data_range(
        worksheet,
        header_row + 1,
        max(
            header_row + len(inconsistent_rows),
            header_row + 1,
        ),
        1,
        max(
            len(dynamic_headers),
            2,
        ),
    )

    for column_index in range(
        1,
        max(
            len(dynamic_headers),
            2,
        ) + 1,
    ):
        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = (
            20
            if column_index <= 2
            else 34
        )

    worksheet.freeze_panes = "A4"


# =========================================================
# 工作表：标签字段矩阵
# =========================================================

def build_field_matrix_sheet(
    workbook: Workbook,
    field_matrix_rows: list[dict[str, Any]],
) -> None:
    """
    创建“标签字段矩阵”工作表。
    """
    worksheet = workbook.create_sheet(
        "标签字段矩阵"
    )

    if field_matrix_rows:
        headers = list(
            field_matrix_rows[0].keys()
        )

    else:
        headers = [
            "标签序号",
            "标签文件名",
        ]

    write_title(
        worksheet,
        "所有标签结构化字段矩阵",
        len(headers),
    )

    header_row = 3

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
    )

    for row_index, matrix_row in enumerate(
        field_matrix_rows,
        start=header_row + 1,
    ):
        for column_index, header in enumerate(
            headers,
            start=1,
        ):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=matrix_row.get(
                    header,
                    "",
                ),
            )

    style_data_range(
        worksheet,
        header_row + 1,
        header_row + len(field_matrix_rows),
        1,
        len(headers),
    )

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = (
            36
            if header == "标签文件名"
            else 20
        )

    worksheet.freeze_panes = "A4"


# =========================================================
# 工作表：OCR明细
# =========================================================

def build_ocr_sheet(
    workbook: Workbook,
    label_results: list[dict[str, Any]],
) -> None:
    """
    创建“OCR识别明细”工作表。
    """
    worksheet = workbook.create_sheet(
        "OCR识别明细"
    )

    headers = [
        "标签序号",
        "标签文件名",
        "OCR序号",
        "识别文字",
        "置信度",
    ]

    write_title(
        worksheet,
        "所有标签OCR识别明细",
        len(headers),
    )

    header_row = 3

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

    style_header_row(
        worksheet,
        header_row,
    )

    current_row = header_row + 1

    for label_index, item in enumerate(
        label_results,
        start=1,
    ):
        filename = safe_text(
            item.get(
                "文件名",
                "",
            )
        )

        ocr_rows = item.get(
            "OCR明细",
            [],
        )

        for ocr_index, ocr_row in enumerate(
            ocr_rows,
            start=1,
        ):
            recognized_text = (
                ocr_row.get(
                    "识别文字",
                    ocr_row.get(
                        "文字",
                        ocr_row.get(
                            "text",
                            "",
                        ),
                    ),
                )
            )

            confidence = (
                ocr_row.get(
                    "置信度",
                    ocr_row.get(
                        "score",
                        "",
                    ),
                )
            )

            values = [
                label_index,
                filename,
                ocr_index,
                recognized_text,
                confidence,
            ]

            for column_index, value in enumerate(
                values,
                start=1,
            ):
                worksheet.cell(
                    row=current_row,
                    column=column_index,
                    value=value,
                )

            current_row += 1

    style_data_range(
        worksheet,
        header_row + 1,
        current_row - 1,
        1,
        len(headers),
    )

    set_column_widths(
        worksheet,
        {
            1: 12,
            2: 36,
            3: 12,
            4: 48,
            5: 16,
        },
    )

    worksheet.freeze_panes = "A4"


# =========================================================
# 主入口
# =========================================================

def generate_batch_excel_report(
    drawing_filename: str,
    batch_overall_result: str,
    label_results: list[dict[str, Any]],
    batch_summary_rows: list[dict[str, Any]],
    batch_detail_rows: list[dict[str, Any]],
    field_matrix_rows: list[dict[str, Any]],
    inconsistent_rows: list[dict[str, Any]],
) -> tuple[bytes, str]:
    """
    生成批量检测总Excel报告。

    返回：
        report_bytes：
            Excel文件的二进制内容。

        report_filename：
            建议下载文件名。
    """
    detection_time = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    workbook = Workbook()

    build_batch_overview_sheet(
        workbook=workbook,
        drawing_filename=drawing_filename,
        batch_overall_result=batch_overall_result,
        label_results=label_results,
        batch_summary_rows=batch_summary_rows,
        inconsistent_rows=inconsistent_rows,
        detection_time=detection_time,
    )

    build_drawing_comparison_sheet(
        workbook=workbook,
        label_results=label_results,
    )

    build_label_horizontal_sheet(
        workbook=workbook,
        batch_detail_rows=batch_detail_rows,
    )

    build_inconsistent_sheet(
        workbook=workbook,
        inconsistent_rows=inconsistent_rows,
    )

    build_field_matrix_sheet(
        workbook=workbook,
        field_matrix_rows=field_matrix_rows,
    )

    build_ocr_sheet(
        workbook=workbook,
        label_results=label_results,
    )

    output_buffer = BytesIO()

    workbook.save(
        output_buffer
    )

    output_buffer.seek(
        0
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    report_filename = (
        f"批量检测总报告_{timestamp}_"
        f"{batch_overall_result}.xlsx"
    )

    return (
        output_buffer.getvalue(),
        report_filename,
    )