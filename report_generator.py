from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# =========================================================
# 样式定义
# =========================================================

TITLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="5B9BD5",
)

PASS_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

FAIL_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

UNKNOWN_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
)

TITLE_FONT = Font(
    color="FFFFFF",
    bold=True,
    size=16,
)

SECTION_FONT = Font(
    bold=True,
    size=12,
)

NORMAL_FONT = Font(
    size=10,
)

PASS_FONT = Font(
    color="274E13",
    bold=True,
)

FAIL_FONT = Font(
    color="9C0006",
    bold=True,
)

UNKNOWN_FONT = Font(
    color="7F6000",
    bold=True,
)

THIN_SIDE = Side(
    style="thin",
    color="B7B7B7",
)

TABLE_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)


# =========================================================
# 通用辅助函数
# =========================================================

def safe_text(value: Any) -> str:
    """
    将任意内容安全转换为字符串。
    """
    if value is None:
        return ""

    return str(value).strip()


def auto_adjust_column_width(
    worksheet,
    min_width: int = 12,
    max_width: int = 45,
) -> None:
    """
    根据单元格内容自动调整Excel列宽。
    """
    for column_cells in worksheet.columns:
        max_length = 0

        column_number = column_cells[0].column
        column_letter = get_column_letter(
            column_number
        )

        for cell in column_cells:
            cell_value = safe_text(
                cell.value
            )

            # 中文显示宽度通常大于英文，
            # 这里简单按字符数乘以系数处理
            display_length = 0

            for character in cell_value:
                if ord(character) > 127:
                    display_length += 2
                else:
                    display_length += 1

            max_length = max(
                max_length,
                display_length,
            )

        adjusted_width = min(
            max(
                max_length + 3,
                min_width,
            ),
            max_width,
        )

        worksheet.column_dimensions[
            column_letter
        ].width = adjusted_width


def apply_table_cell_style(
    cell,
) -> None:
    """
    设置普通表格单元格样式。
    """
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )
    cell.font = NORMAL_FONT


def apply_header_style(
    cell,
) -> None:
    """
    设置表头样式。
    """
    cell.fill = HEADER_FILL
    cell.font = WHITE_FONT
    cell.border = TABLE_BORDER
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


def apply_result_style(
    cell,
    result: str,
) -> None:
    """
    根据PASS、FAIL或无法判断设置颜色。
    """
    normalized_result = safe_text(
        result
    ).upper()

    if normalized_result == "PASS":
        cell.fill = PASS_FILL
        cell.font = PASS_FONT

    elif normalized_result == "FAIL":
        cell.fill = FAIL_FILL
        cell.font = FAIL_FONT

    else:
        cell.fill = UNKNOWN_FILL
        cell.font = UNKNOWN_FONT


# =========================================================
# 报告工作表
# =========================================================

def create_summary_sheet(
    workbook: Workbook,
    report_info: dict,
    comparison_rows: list[dict],
) -> None:
    """
    创建检测报告主表。
    """
    worksheet = workbook.active
    worksheet.title = "检测报告"

    worksheet.freeze_panes = "A10"

    # -----------------------------------------------------
    # 报告标题
    # -----------------------------------------------------

    worksheet.merge_cells(
        "A1:E2"
    )

    title_cell = worksheet["A1"]
    title_cell.value = (
        "制造业图纸与标签自动一致性检测报告"
    )
    title_cell.fill = TITLE_FILL
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 24
    worksheet.row_dimensions[2].height = 12

    # -----------------------------------------------------
    # 基础信息
    # -----------------------------------------------------

    worksheet.merge_cells(
        "A4:E4"
    )

    section_cell = worksheet["A4"]
    section_cell.value = "一、检测基本信息"
    section_cell.fill = SECTION_FILL
    section_cell.font = SECTION_FONT
    section_cell.alignment = Alignment(
        vertical="center",
    )

    base_info_rows = [
        (
            "检测编号",
            report_info.get(
                "检测编号",
                "",
            ),
            "检测时间",
            report_info.get(
                "检测时间",
                "",
            ),
        ),
        (
            "图纸文件",
            report_info.get(
                "图纸文件",
                "",
            ),
            "标签文件",
            report_info.get(
                "标签文件",
                "",
            ),
        ),
        (
            "总体结果",
            report_info.get(
                "总体结果",
                "",
            ),
            "一致字段数量",
            report_info.get(
                "一致字段数量",
                0,
            ),
        ),
        (
            "异常字段数量",
            report_info.get(
                "异常字段数量",
                0,
            ),
            "无法判断数量",
            report_info.get(
                "无法判断数量",
                0,
            ),
        ),
    ]

    start_row = 5

    for row_offset, row_data in enumerate(
        base_info_rows
    ):
        row_number = (
            start_row + row_offset
        )

        worksheet.cell(
            row=row_number,
            column=1,
            value=row_data[0],
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=row_data[1],
        )

        worksheet.cell(
            row=row_number,
            column=3,
            value=row_data[2],
        )

        worksheet.merge_cells(
            start_row=row_number,
            start_column=4,
            end_row=row_number,
            end_column=5,
        )

        worksheet.cell(
            row=row_number,
            column=4,
            value=row_data[3],
        )

        for column_number in range(
            1,
            6,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            apply_table_cell_style(
                cell
            )

        worksheet.cell(
            row=row_number,
            column=1,
        ).font = Font(
            bold=True
        )

        worksheet.cell(
            row=row_number,
            column=3,
        ).font = Font(
            bold=True
        )

    overall_result_cell = worksheet[
        "B7"
    ]

    apply_result_style(
        overall_result_cell,
        report_info.get(
            "总体结果",
            "",
        ),
    )

    # -----------------------------------------------------
    # 比对明细
    # -----------------------------------------------------

    worksheet.merge_cells(
        "A10:E10"
    )

    comparison_section = worksheet[
        "A10"
    ]

    comparison_section.value = (
        "二、图纸与标签字段比对明细"
    )
    comparison_section.fill = SECTION_FILL
    comparison_section.font = SECTION_FONT

    header_row = 11

    headers = [
        "字段名称",
        "图纸值",
        "标签值",
        "检测结果",
        "异常说明",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_number,
            value=header,
        )

        apply_header_style(
            cell
        )

    for row_offset, item in enumerate(
        comparison_rows,
        start=1,
    ):
        row_number = (
            header_row + row_offset
        )

        values = [
            item.get(
                "字段名称",
                "",
            ),
            item.get(
                "图纸值",
                "",
            ),
            item.get(
                "标签值",
                "",
            ),
            item.get(
                "检测结果",
                "",
            ),
            item.get(
                "异常说明",
                "",
            ),
        ]

        for column_number, value in enumerate(
            values,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

            apply_table_cell_style(
                cell
            )

        result = safe_text(
            item.get(
                "检测结果",
                "",
            )
        )

        for column_number in range(
            1,
            6,
        ):
            apply_result_style(
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ),
                result,
            )

    worksheet.auto_filter.ref = (
        f"A{header_row}:"
        f"E{header_row + len(comparison_rows)}"
    )

    worksheet.page_setup.orientation = (
        "landscape"
    )

    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    auto_adjust_column_width(
        worksheet
    )


# =========================================================
# OCR明细工作表
# =========================================================

def create_ocr_detail_sheet(
    workbook: Workbook,
    ocr_rows: list[dict],
) -> None:
    """
    创建OCR识别明细工作表。
    """
    worksheet = workbook.create_sheet(
        "OCR识别明细"
    )

    headers = [
        "序号",
        "识别文字",
        "置信度",
        "置信度百分比",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column_number,
            value=header,
        )

        apply_header_style(
            cell
        )

    for row_number, item in enumerate(
        ocr_rows,
        start=2,
    ):
        score = item.get(
            "置信度",
            0,
        )

        try:
            numeric_score = float(
                score
            )
        except (
            TypeError,
            ValueError,
        ):
            numeric_score = 0.0

        values = [
            item.get(
                "序号",
                row_number - 1,
            ),
            item.get(
                "识别文字",
                "",
            ),
            numeric_score,
            numeric_score,
        ]

        for column_number, value in enumerate(
            values,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

            apply_table_cell_style(
                cell
            )

        worksheet.cell(
            row=row_number,
            column=3,
        ).number_format = "0.0000"

        worksheet.cell(
            row=row_number,
            column=4,
        ).number_format = "0.00%"

        if numeric_score < 0.80:
            for column_number in range(
                1,
                5,
            ):
                cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                )

                cell.fill = UNKNOWN_FILL
                cell.font = UNKNOWN_FONT

    worksheet.freeze_panes = "A2"

    if ocr_rows:
        worksheet.auto_filter.ref = (
            f"A1:D{len(ocr_rows) + 1}"
        )

    auto_adjust_column_width(
        worksheet
    )


# =========================================================
# 原始文字工作表
# =========================================================

def create_raw_text_sheet(
    workbook: Workbook,
    drawing_text: str,
    label_text: str,
) -> None:
    """
    保存PDF原始文字和标签OCR原始文字。
    """
    worksheet = workbook.create_sheet(
        "原始文字"
    )

    worksheet["A1"] = "PDF图纸原始文字"
    worksheet["A1"].fill = SECTION_FILL
    worksheet["A1"].font = SECTION_FONT

    worksheet["A2"] = drawing_text
    worksheet["A2"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    worksheet["B1"] = "标签OCR原始文字"
    worksheet["B1"].fill = SECTION_FILL
    worksheet["B1"].font = SECTION_FONT

    worksheet["B2"] = label_text
    worksheet["B2"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    worksheet.column_dimensions[
        "A"
    ].width = 60

    worksheet.column_dimensions[
        "B"
    ].width = 60

    worksheet.row_dimensions[
        2
    ].height = 420


# =========================================================
# 检测报告生成入口
# =========================================================

def generate_excel_report(
    drawing_filename: str,
    label_filename: str,
    overall_result: str,
    comparison_rows: list[dict],
    ocr_rows: list[dict],
    drawing_text: str,
    label_text: str,
) -> tuple[bytes, str]:
    """
    生成完整Excel检测报告。

    返回：
        report_bytes：
            Excel文件的二进制内容。

        report_filename：
            自动生成的Excel文件名。
    """
    current_time = datetime.now()

    detection_id = current_time.strftime(
        "OCR-%Y%m%d-%H%M%S"
    )

    detection_time = current_time.strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    pass_count = sum(
        1
        for row in comparison_rows
        if row.get(
            "检测结果"
        ) == "PASS"
    )

    fail_count = sum(
        1
        for row in comparison_rows
        if row.get(
            "检测结果"
        ) == "FAIL"
    )

    unknown_count = sum(
        1
        for row in comparison_rows
        if row.get(
            "检测结果"
        ) == "无法判断"
    )

    report_info = {
        "检测编号": detection_id,
        "检测时间": detection_time,
        "图纸文件": drawing_filename,
        "标签文件": label_filename,
        "总体结果": overall_result,
        "一致字段数量": pass_count,
        "异常字段数量": fail_count,
        "无法判断数量": unknown_count,
    }

    workbook = Workbook()

    create_summary_sheet(
        workbook=workbook,
        report_info=report_info,
        comparison_rows=comparison_rows,
    )

    create_ocr_detail_sheet(
        workbook=workbook,
        ocr_rows=ocr_rows,
    )

    create_raw_text_sheet(
        workbook=workbook,
        drawing_text=drawing_text,
        label_text=label_text,
    )

    output_stream = BytesIO()

    workbook.save(
        output_stream
    )

    output_stream.seek(0)

    report_filename = (
        f"检测报告_"
        f"{current_time.strftime('%Y%m%d_%H%M%S')}_"
        f"{overall_result}.xlsx"
    )

    return (
        output_stream.getvalue(),
        report_filename,
    )